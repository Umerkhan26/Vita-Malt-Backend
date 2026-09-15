"""
Convert Vita Malt crown Excel → clean codes CSV for batched import.

Only keeps GRAND PRIZE ENTRY rows (skips INSTANT WIN).

Usage:
  python scripts/xlsxToCodesCsv.py "C:/Users/SUN RISE/Downloads/VitaMalt_UTC_Crown_Codes_2026 - Copy.xlsx"

Creates:
  backend/data/crown-codes-grand-prize.csv   (one code per line)
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl first:  pip install openpyxl")
    sys.exit(1)


def main() -> None:
    if len(sys.argv) < 2:
        print('Usage: python scripts/xlsxToCodesCsv.py "<path-to.xlsx>"')
        sys.exit(1)

    xlsx = Path(sys.argv[1]).expanduser().resolve()
    if not xlsx.exists():
        print(f"File not found: {xlsx}")
        sys.exit(1)

    out_dir = Path(__file__).resolve().parent.parent / "data"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "crown-codes-grand-prize.csv"

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheet_name = "Master List - Supplier" if "Master List - Supplier" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    codes: list[str] = []
    instant = 0
    seen: set[str] = set()

    # Rows: title, total, blank, header, then data
    for row in ws.iter_rows(min_row=5, values_only=True):
        print_text = str(row[1] or "").strip().upper()
        kind = str(row[2] or "").strip().upper()
        if "INSTANT" in kind:
            instant += 1
            continue
        if not print_text or len(print_text) < 4:
            continue
        # skip header leftovers / prize labels
        if print_text in {"PRINTTEXT", "PRINT TEXT", "PRINTTEXT/CODE", "MERCH", "CODE", "CODES"}:
            continue
        if not all(c.isalnum() or c == "-" for c in print_text):
            continue
        if print_text in seen:
            continue
        seen.add(print_text)
        codes.append(print_text)

    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["code"])
        for c in codes:
            w.writerow([c])

    print(f"Sheet: {sheet_name}")
    print(f"Instant win rows skipped: {instant}")
    print(f"Grand prize codes written: {len(codes)}")
    print(f"Output: {out}")
    print()
    print("Next — import in STEPS (example 5,000 per run):")
    print(f'  cd backend')
    print(f'  npm run import-csv -- "{out}" --batch-size 5000 --once')
    print("  (run the same command again for the next chunk)")


if __name__ == "__main__":
    main()
